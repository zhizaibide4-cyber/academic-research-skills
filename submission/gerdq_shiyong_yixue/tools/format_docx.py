"""Apply Chinese-journal formatting to the pandoc-generated manuscript .docx.

Body: 宋体 / Times New Roman 10.5 pt (五号), 1.5 line spacing; headings 黑体;
tables as three-line tables (三线表); 【作者…】 placeholders highlighted.
Usage: python3 format_docx.py in.docx out.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

from docx.shared import Cm, Pt, RGBColor

src, out = sys.argv[1], sys.argv[2]
doc = Document(src)

for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.8)


def set_font(run, east="宋体", west="Times New Roman", size=None, bold=None):
    run.font.name = west
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), west)
    rfonts.set(qn("w:eastAsia"), east)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in [s.name for s in doc.styles]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        st.element.get_or_add_rPr()
        rf = st.element.rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            st.element.rPr.insert(0, rf)
        rf.set(qn("w:eastAsia"), "宋体")
        rf.set(qn("w:ascii"), "Times New Roman")
        rf.set(qn("w:hAnsi"), "Times New Roman")

HEAD_SIZE = {"Title": 16, "Heading 1": 16, "Heading 2": 12, "Heading 3": 10.5}
PLACEHOLDER = re.compile(r"【作者[^】]*】|【 】")

in_refs = False
for p in doc.paragraphs:
    style = p.style.name
    text = p.text
    if style in HEAD_SIZE:
        for r in p.runs:
            set_font(r, east="黑体", west="Times New Roman", size=HEAD_SIZE[style], bold=True)
            r.font.color.rgb = RGBColor(0, 0, 0)
        if style in ("Title", "Heading 1"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12 if style != "Heading 3" else 6)
        p.paragraph_format.space_after = Pt(6)
        in_refs = text.strip() == "参考文献"
        continue
    is_ref = in_refs and re.match(r"^\[\d+\]", text)
    size = 9 if is_ref else 10.5
    for r in p.runs:
        set_font(r, size=size)
    pf = p.paragraph_format
    pf.line_spacing = 1.25 if is_ref else 1.5
    pf.space_after = Pt(0 if is_ref else 3)
    if is_ref:
        pf.left_indent = Cm(0.75)
        pf.first_line_indent = Cm(-0.75)
    elif style in ("Body Text", "First Paragraph", "Normal") and len(text) > 60 and not text.startswith(("【", "¹", "**")):
        pf.first_line_indent = Pt(21)
    if text.startswith(("路源", "LU Yuan", "¹", "通信作者", "Corresponding", "基金项目")):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = None
        for r in p.runs:
            r.font.size = Pt(9)
    if re.match(r"^\*?\*?表\d", text) or text.startswith("表") and "　" in text[:4]:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = None
        for r in p.runs:
            set_font(r, east="黑体", size=9, bold=True)
    if text.startswith("注："):
        pf.first_line_indent = None
        for r in p.runs:
            r.font.size = Pt(9)



def add_highlight(r_el):
    rpr = r_el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r_el.insert(0, rpr)
    if rpr.find(qn("w:highlight")) is None:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        rpr.append(hl)


def highlight_placeholders(paragraph):
    if "【" not in paragraph.text:
        return
    # Placeholders may span several runs (e.g. around sub/superscript), so
    # split runs at bracket boundaries first, then highlight between them.
    for run in list(paragraph.runs):
        parts = [x for x in re.split(r"(【|】)", run.text) if x]
        if len(parts) < 2:
            continue
        anchor = run._element
        for part in parts:
            new = deepcopy(run._element)
            for t in new.findall(qn("w:t")):
                new.remove(t)
            t = OxmlElement("w:t")
            t.text = part
            t.set(qn("xml:space"), "preserve")
            new.append(t)
            anchor.addnext(new)
            anchor = new
        run._element.getparent().remove(run._element)
    runs = paragraph.runs
    inside = False
    for i, run in enumerate(runs):
        if run.text == "【":
            following = "".join(r.text for r in runs[i + 1:i + 4])
            inside = following.startswith("作者") or following.startswith(" 】")
        if inside:
            add_highlight(run._element)
        if run.text == "】":
            inside = False


for p in doc.paragraphs:
    highlight_placeholders(p)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                highlight_placeholders(p)

for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    if name in [x.name for x in doc.styles]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)


def border(tag, sz):
    b = OxmlElement(f"w:{tag}")
    b.set(qn("w:val"), "single" if sz else "nil")
    if sz:
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
    return b


for t in doc.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tb = OxmlElement("w:tblBorders")
    for tag, sz in (("top", 12), ("left", 0), ("bottom", 12), ("right", 0), ("insideH", 0), ("insideV", 0)):
        tb.append(border(tag, sz))
    tblPr.append(tb)
    for old in tblPr.findall(qn("w:tblStyle")):
        tblPr.remove(old)
    ncol = len(t.columns)
    def text_len(s, cjk):
        return sum(cjk if ord(ch) > 0x2E80 else 1 for ch in s)

    lens = [max([text_len(t.rows[0].cells[j].text, 1.2)]
                + [text_len(row.cells[j].text, 2) for row in t.rows[1:]])
            for j in range(ncol)]
    weights = [min(max(l, 9), 28) + 3 for l in lens]
    weights[0] = min(max(lens[0], 16), 20 if ncol >= 6 else 28)
    total_cm = 15.4
    widths = [Cm(total_cm * w / sum(weights)) for w in weights]
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(total_cm / 2.54 * 1440)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j].cm / 2.54 * 1440)))
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            if i == 0:
                tcb = OxmlElement("w:tcBorders")
                tcb.append(border("bottom", 6))
                tcPr.append(tcb)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cell is row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_font(r, size=9, bold=True if i == 0 else None)

doc.save(out)
print("saved", out)
