"""Step 2: build every manuscript table (表1~表7) and supplementary table (附表S1~S6).

Each table is written as  output/tables/<名称>.csv  (open in Excel) and  .md  (inserted into the
manuscript by tools/build.py), and all tables are collected in  output/tables/统计表汇总.xlsx.

Values marked as recomputed are calculated here from counts; continuous-variable descriptions and
rank-test statistics (H, U) are copied from the original manuscript because no individual data exist.
"""

import csv
import math

from openpyxl import Workbook
from scipy.stats import fisher_exact, norm
from openpyxl.styles import Alignment, Border, Font, Side

from common import (CHK, TAB, Z, ci_pct, chi2_rxc, cochran_armitage, cohen_kappa, e_value_or_common,
                    fmt_p, p_eq, grouped_auc_delong, hanley_mcneil_se, hedges_g, likelihood_ratio, load,
                    load_reconstructed, msd, odds_ratio_woolf, pct, post_test, wilson, write_table)

D = load()
R = load_reconstructed()
N = D["N"]
tables = []

# ---------------------------------------------------------------------------
# 表1
# ---------------------------------------------------------------------------
g = D["groups"]
hdr = ["项目", "轻度肥胖组（*n*=63）", "中度肥胖组（*n*=26）", "重度肥胖组（*n*=14）", "合计（*n*=103）", "统计量", "*P*值"]
rows = []
for key, label, d in [("age", "年龄（岁，$\\bar{x}$±*s*）", 2), ("height", "身高（m，$\\bar{x}$±*s*）", 2),
                      ("weight", "体重（kg，$\\bar{x}$±*s*）", 2), ("bmi", "BMI（kg/m²，$\\bar{x}$±*s*）", 2)]:
    rows.append([label] + [msd(m, s, d) for m, s in zip(g[f"{key}_mean"], g[f"{key}_sd"])]
                + [msd(*g[f"{key}_total"], d), f"*H*={g[f'{key}_H']:.2f}", fmt_p(g[f"{key}_P"])])
TOT = dict(male=55, smoke_drink=34, endo_abnormal=39, ppi_test_pos=57, gerd=79)
for key, label in [("male", "男性［例（%）］"), ("smoke_drink", "吸烟和（或）饮酒［例（%）］"),
                   ("endo_abnormal", "胃镜异常［例（%）］"), ("ppi_test_pos", "PPI试验阳性［例（%）］"),
                   ("gerd", "GERD复合诊断［例（%）］")]:
    k = g[key]
    c2, p, _ = chi2_rxc([k, [n - x for n, x in zip(g["n"], k)]])
    rows.append([label] + [pct(x, n) for x, n in zip(k, g["n"])] + [pct(TOT[key], N), f"χ²={c2:.2f}", fmt_p(p)])
mq = g["gerdq_median_q1_q3"]
fm = lambda v: f"{v[0]:g}（{v[1]:g}，{v[2]:g}）"
rows.append(["基线GerdQ总分［分，*M*（*Q*₁，*Q*₃）］"] + [fm(v) for v in mq] + [fm(g["gerdq_total_median_q1_q3"]),
             f"*H*={g['gerdq_H']:.2f}", fmt_p(g["gerdq_P"])])
rows.append(["基线4GerdQ（分，$\\bar{x}$±*s*）"] + [msd(m, s) for m, s in zip(g["gerdq4_mean"], g["gerdq4_sd"])]
            + [msd(*g["gerdq4_total"]), f"*H*={g['gerdq4_H']:.2f}", fmt_p(g["gerdq4_P"])])
rows.append(["减重幅度（kg，$\\bar{x}$±*s*）"] + [msd(m, s) for m, s in zip(g["wl_mean"], g["wl_sd"])]
            + [msd(*g["wl_total"]), f"*H*={g['wl_H']:.2f}", fmt_p(g["wl_P"])])
tables.append(write_table("表1_不同肥胖程度组基线资料", hdr, rows,
    "注：BMI为体质指数；PPI为质子泵抑制剂；GERD为胃食管反流病；4GerdQ为GerdQ第1、2、5、6条得分之和；减重幅度为减重管理后体重较基线的下降值。"
    "计数资料组间比较采用χ²检验；计量资料采用Kruskal-Wallis *H*检验"))

# ---------------------------------------------------------------------------
# 表2
# ---------------------------------------------------------------------------
it = D["items"]
hdr = ["条目", "0分［例（%）］", "1分［例（%）］", "2分［例（%）］", "3分［例（%）］", "得分（分，$\\bar{x}$±*s*）", "条目-总分*r*ₛ"]
rows = []
for nm in it["order"]:
    cnt = it["counts"][nm]
    mean = sum(i * c for i, c in enumerate(cnt)) / N
    star = "\\*" if nm in it["negative_items"] else ""
    rows.append([nm + star] + [pct(c, N) for c in cnt] + [msd(mean, R["item_sd"][nm]), f"{it['item_total_rs'][nm]:.3f}"])
tables.append(write_table("表2_基线GerdQ条目得分分布", hdr, rows,
    "注：\\*阴性预测条目，反向计分（得分越高表示症状天数越少）；条目-总分相关采用Spearman秩相关"))

# ---------------------------------------------------------------------------
# 表3
# ---------------------------------------------------------------------------
dg = D["diagnosis"]
npos, nneg = dg["n_pos"], dg["n_neg"]
hdr = ["截断值（分）", "灵敏度［%（95%CI）］", "特异度［%（95%CI）］", "阳性预测值（%）", "阴性预测值（%）", "约登指数"]
rows = []
cut_rows = []
lr_notes = []
for c in dg["cutoffs"]:
    tp, fp = R["pos_ge"][str(c)], R["neg_ge"][str(c)]
    fn, tn = npos - tp, nneg - fp
    se_lo, se_hi = wilson(tp, npos)
    sp_lo, sp_hi = wilson(tn, nneg)
    lrp = (tp / npos) / (fp / nneg)
    lrn = (fn / npos) / (tn / nneg)
    rows.append([f"≥{c}", f"{100 * tp / npos:.1f}（{100 * se_lo:.1f}~{100 * se_hi:.1f}）",
                 f"{100 * tn / nneg:.1f}（{100 * sp_lo:.1f}~{100 * sp_hi:.1f}）",
                 f"{100 * tp / (tp + fp):.1f}", f"{100 * tn / (tn + fn):.1f}", f"{tp / npos + tn / nneg - 1:.3f}"])
    lr_notes.append(f"≥{c}分{lrp:.2f}/{lrn:.2f}")
    cut_rows.append([c, tp, fp, fn, tn])
tables.append(write_table("表3_GerdQ不同截断值的诊断效能", hdr, rows,
    "注：GERD复合诊断阳性79例、阴性24例；灵敏度和特异度的95%CI采用Wilson法；阳性/阴性似然比分别为" + "，".join(lr_notes)
    + "；因入组要求门诊GerdQ≥7分，本表不能反映≤7分时的诊断效能"))

# ---------------------------------------------------------------------------
# 表4
# ---------------------------------------------------------------------------
pos_b, neg_b = R["score_pos"], R["score_neg"]
merged = [("7分", [0]), ("8~9分", [1, 2]), ("≥10分", [3, 4, 5])]
pre = npos / N
hdr = ["总分区间", "阳性/阴性（例）", "分层似然比（95%CI）", "本组验后概率（%）", "验后概率^a^（%）", "验后概率^b^（%）"]
rows = []
sslr = []
for lab, idx in merged:
    p = sum(pos_b[i] for i in idx)
    n = sum(neg_b[i] for i in idx)
    lr, lo, hi = likelihood_ratio(p, npos, n, nneg)
    sslr.append((lab, p, n, lr, lo, hi))
    rows.append([lab, f"{p}/{n}", f"{lr:.2f}（{lo:.2f}~{hi:.2f}）", f"{100 * p / (p + n):.1f}",
                 f"{100 * post_test(0.30, lr):.1f}", f"{100 * post_test(0.50, lr):.1f}"])
tables.append(write_table("表4_GerdQ分层似然比及验后概率", hdr, rows,
    f"注：^a^验前概率为30%时；^b^验前概率为50%时；各区间例数为GERD复合诊断阳性/阴性例数，由表3截断值结果反推；"
    f"本组验前概率（GERD复合诊断比例）为{100 * pre:.1f}%；分层似然比的95%CI采用对数法；"
    "验后概率=验后比值/（1+验后比值），验后比值=验前比值×分层似然比"))

# ---------------------------------------------------------------------------
# 表5
# ---------------------------------------------------------------------------
t = D["treatment"]
lo, hi = wilson(t["improved"], N)
lo8, hi8 = wilson(t["post_le1"], N)
hdr = ["指标", "结果", "统计量或95%CI"]
rows = [
    ["治疗前4GerdQ（分，$\\bar{x}$±*s*）", msd(*t["gerdq4_pre"]), "—"],
    ["治疗后4GerdQ（分，$\\bar{x}$±*s*）", msd(*t["gerdq4_post"]), "—"],
    ["4GerdQ下降值（分，$\\bar{x}$±*s*）", msd(*t["delta"]),
     f"95%CI：{t['delta_ci'][0]:.2f}~{t['delta_ci'][1]:.2f}；Wilcoxon *W*={t['wilcoxon_W']}，*P*<0.001；配对*t*={t['paired_t']:.2f}，*P*<0.001"],
    ["症状改善（下降≥1分）［例（%）］", pct(t["improved"], N), f"95%CI：{ci_pct(lo, hi)}"],
    ["改善者4GerdQ下降值（分，$\\bar{x}$±*s*）", msd(R["responder_delta_mean"], R["responder_delta_sd"]), "—"],
    ["治疗后4GerdQ≤1分［例（%）］", pct(t["post_le1"], N), f"95%CI：{ci_pct(lo8, hi8)}"],
    ["GerdQ总分（治疗前→治疗后，分，$\\bar{x}$±*s*）", f"{msd(*t['total_pre'])}→{msd(*t['total_post'])}", "Wilcoxon *P*<0.001"],
    ["上腹痛+恶心条目合计（治疗前→治疗后，分）",
     f"{t['total_pre'][0] - t['gerdq4_pre'][0]:.2f}→{t['total_post'][0] - t['gerdq4_post'][0]:.2f}", "—"],
    ["减重幅度（kg，$\\bar{x}$±*s*）", msd(*t["wl"]), f"范围：{t['wl_range'][0]}~{t['wl_range'][1]}"],
]
tables.append(write_table("表5_治疗前后GerdQ评分及体重变化", hdr, rows,
    "注：4GerdQ为GerdQ第1、2、5、6条得分之和；下降值=治疗前−治疗后；率的95%CI采用Wilson法；"
    "改善者下降值由全体下降值的均数、标准差及未改善者下降值均为0推算；上腹痛+恶心条目合计=GerdQ总分−4GerdQ"))

# ---------------------------------------------------------------------------
# 表6
# ---------------------------------------------------------------------------
def group_thousands(s):
    whole, _, frac = s.partition(".")
    whole = f"{int(whole):,}".replace(",", " ")
    return whole + ("." + frac if frac else "")


r = D["response"]
hdr = ["项目", "改善组（*n*=72）", "未改善组（*n*=31）", "统计量", "*P*值"]
rows = []
for key, label in [("age", "年龄（岁，$\\bar{x}$±*s*）"), ("bmi", "BMI（kg/m²，$\\bar{x}$±*s*）"),
                   ("weight", "体重（kg，$\\bar{x}$±*s*）"), ("gerdq4", "基线4GerdQ（分，$\\bar{x}$±*s*）"),
                   ("gerdq_total", "基线GerdQ总分（分，$\\bar{x}$±*s*）"), ("wl", "减重幅度（kg，$\\bar{x}$±*s*）"),
                   ("wl_pct", "减重百分比（%，$\\bar{x}$±*s*）")]:
    a, b, u, p = r[key]
    rows.append([label, msd(*a), msd(*b), "*U*=" + group_thousands(u[2:]), fmt_p(p)])
for key, label in [("male", "男性［例（%）］"), ("smoke_drink", "吸烟和（或）饮酒［例（%）］"),
                   ("endo_abnormal", "胃镜异常［例（%）］"), ("gerd", "GERD复合诊断［例（%）］")]:
    a, b = r[key][:2]
    c2, p, _ = chi2_rxc([[a, 72 - a], [b, 31 - b]], yates=True)
    rows.append([label, pct(a, 72), pct(b, 31), f"χ²={c2:.2f}", fmt_p(p)])
tables.append(write_table("表6_症状改善组与未改善组比较", hdr, rows,
    "注：症状改善定义为治疗后4GerdQ较基线下降≥1分；计量资料采用Mann-Whitney *U*检验；计数资料采用χ²检验（2×2表连续性校正）"))

# ---------------------------------------------------------------------------
# 表7
# ---------------------------------------------------------------------------
lg = D["logistic"]
hdr = ["变量", "*β*", "*SE*", "Wald χ²", "*OR*（95%CI）", "*P*值"]
rows = []
for nm, (b, se, *_rest) in lg["terms"].items():
    wald = (b / se) ** 2
    p = 2 * norm.sf(abs(b / se))
    orci = "—" if nm == "常数项" else f"{math.exp(b):.2f}（{math.exp(b - Z * se):.2f}~{math.exp(b + Z * se):.2f}）"
    rows.append([nm, f"{b:.3f}".replace("-", "−"), f"{se:.3f}", f"{wald:.2f}", orci, fmt_p(p)])
epv = r["n"][1] / lg["candidates"]
tables.append(write_table("表7_症状改善相关因素的多因素Logistic回归", hdr, rows,
    f"注：候选变量为年龄、性别、BMI、吸烟和（或）饮酒、高风险职业、基线4GerdQ、胃镜结果和减重幅度，后退法（*P*>0.10剔除）；"
    f"Wald χ²、*OR*及95%CI由*β*和*SE*重新计算；未改善组31例，事件数与候选变量数之比约为{epv:.1f}，结果存在过拟合风险"))

# ---------------------------------------------------------------------------
# 附表S1 重建评分分布
# ---------------------------------------------------------------------------
hdr = ["GerdQ总分（分）", "GERD复合诊断阳性（例）", "GERD复合诊断阴性（例）", "合计（例）", "GERD比例（%）", "逐分似然比（95%CI）"]
rows = []
for b, p, n in zip(R["score_bands"], pos_b, neg_b):
    lr, lo, hi = likelihood_ratio(p, npos, n, nneg)
    rows.append([b, p, n, p + n, f"{100 * p / (p + n):.1f}", f"{lr:.2f}（{lo:.2f}~{hi:.2f}）"])
rows.append(["合计", npos, nneg, N, f"{100 * npos / N:.1f}", "—"])
scores = [7, 8, 9, 10, 11, 12]
z_tr, p_tr = cochran_armitage(pos_b, [a + b for a, b in zip(pos_b, neg_b)], scores)
auc_g, se_g = R["auc_grouped"], R["auc_grouped_se"]
tables.append(write_table("附表S1_重建的GerdQ总分与GERD交叉分布", hdr, rows,
    f"注：各分值例数由原稿截断值表的灵敏度、特异度唯一反推，并与原稿PPV、NPV、约登指数、*U*值逐一核对一致；"
    f"GERD比例随分值升高的趋势检验：Cochran-Armitage *Z*={z_tr:.2f}，{p_eq(p_tr)}（≥12分按12分计）；"
    f"按本表计算的AUC={auc_g:.3f}（DeLong法95%CI：{auc_g - Z * se_g:.3f}~{auc_g + Z * se_g:.3f}）"))

# ---------------------------------------------------------------------------
# 附表S2 诊断路径
# ---------------------------------------------------------------------------
names, pn, presp, pge10 = R["pathway_names"], R["pathway_n"], R["pathway_resp"], R["pathway_ge10"]
z10, p10 = cochran_armitage(pge10, pn, [2, 1, 0])
zr, pr = cochran_armitage(presp, pn, [2, 1, 0])
hdr = ["诊断路径", "例数", "GerdQ≥10分［例（%）］", "GerdQ≥10分 95%CI", "4周症状改善［例（%）］", "症状改善 95%CI"]
rows = []
for nm, n, k10, kr in zip(names, pn, pge10, presp):
    a, b = wilson(k10, n)
    c, d = wilson(kr, n)
    rows.append([nm, n, pct(k10, n), ci_pct(a, b), pct(kr, n), ci_pct(c, d)])
rows.append(["趋势检验", "—", f"*Z*={z10:.2f}", f"{p_eq(p10)}", f"*Z*={zr:.2f}", f"{p_eq(pr)}"])
sub = [[pge10[1], pn[1] - pge10[1]], [pge10[2], pn[2] - pge10[2]]]
or_sub, p_sub = fisher_exact(sub)
se_sub = wilson(pge10[1], pn[1])
sp_sub = wilson(pn[2] - pge10[2], pn[2])
tables.append(write_table("附表S2_不同诊断路径的GerdQ评分与疗效", hdr, rows,
    f"注：例数由原稿报告数据推算（胃镜异常39例；胃镜阴性且PPI试验阳性40例；GERD阴性24例；改善组中GERD阳性59例、胃镜异常33例；"
    f"GerdQ≥10分者在两种结局下的灵敏度、特异度）；趋势检验为Cochran-Armitage检验（评分2、1、0）；率的95%CI采用Wilson法。"
    f"胃镜阴性者（64例）中以GerdQ≥10分识别PPI试验阳性：灵敏度{100 * pge10[1] / pn[1]:.1f}%（{ci_pct(*se_sub)}），"
    f"特异度{100 * (pn[2] - pge10[2]) / pn[2]:.1f}%（{ci_pct(*sp_sub)}），*OR*={or_sub:.2f}，Fisher确切概率法{p_eq(p_sub)}"))

# ---------------------------------------------------------------------------
# 附表S3 单因素与多因素OR
# ---------------------------------------------------------------------------
hdr = ["变量", "分析", "*OR*（95%CI）", "*P*值", "E值（点估计）", "E值（95%CI下限）"]
rows = []
uni = []
for key, label in [("male", "男性"), ("smoke_drink", "吸烟和（或）饮酒"), ("endo_abnormal", "胃镜异常"), ("gerd", "GERD复合诊断")]:
    a, b = r[key][:2]
    o, lo, hi = odds_ratio_woolf(a, b, 72 - a, 31 - b)
    _, p, _ = chi2_rxc([[a, 72 - a], [b, 31 - b]], yates=True)
    ev_lo = e_value_or_common(lo) if lo > 1 else 1.0
    rows.append([label, "单因素", f"{o:.2f}（{lo:.2f}~{hi:.2f}）", fmt_p(p), f"{e_value_or_common(o):.2f}", f"{ev_lo:.2f}"])
    uni.append((label, o, lo, hi))
multi = []
for nm, (b, se, *_x) in lg["terms"].items():
    if nm == "常数项":
        continue
    o, lo, hi = math.exp(b), math.exp(b - Z * se), math.exp(b + Z * se)
    p = 2 * norm.sf(abs(b / se))
    rows.append([nm, "多因素", f"{o:.2f}（{lo:.2f}~{hi:.2f}）", fmt_p(p), f"{e_value_or_common(o):.2f}", f"{e_value_or_common(lo):.2f}"])
    multi.append((nm, o, lo, hi))
g_wl = hedges_g(*r["wl"][0], 72, *r["wl"][1], 31)
g_wp = hedges_g(*r["wl_pct"][0], 72, *r["wl_pct"][1], 31)
tables.append(write_table("附表S3_症状改善的单因素与多因素OR及E值", hdr, rows,
    f"注：单因素*OR*由原稿表6四格表计算（行：改善组、未改善组；列：有、无该特征，即*OR*=(a×d)/(b×c)），95%CI采用Woolf法，*P*值为连续性校正χ²检验；"
    f"计量资料无原始数据，无法计算单因素*OR*，以Hedges *g*表示：减重幅度*g*={g_wl:.2f}，减重百分比*g*={g_wp:.2f}；"
    "E值按VanderWeele法计算（结局常见，先以√*OR*近似换算为*RR*），表示能完全解释该关联所需的未测混杂强度（*RR*尺度），95%CI下限≤1时E值记为1.00"))

# ---------------------------------------------------------------------------
# 附表S4 PPI试验与4周疗效一致性
# ---------------------------------------------------------------------------
agree = [[presp[1], pn[1] - presp[1]], [presp[2], pn[2] - presp[2]]]
k, klo, khi, po = cohen_kappa(agree)
hdr = ["胃镜阴性者（*n*=64）", "4周症状改善（例）", "4周未改善（例）", "合计（例）"]
rows = [["PPI试验阳性", agree[0][0], agree[0][1], sum(agree[0])],
        ["PPI试验阴性", agree[1][0], agree[1][1], sum(agree[1])],
        ["合计", agree[0][0] + agree[1][0], agree[0][1] + agree[1][1], 64]]
tables.append(write_table("附表S4_PPI试验与4周症状改善的一致性", hdr, rows,
    f"注：一致率{100 * po:.1f}%；Cohen κ={k:.2f}（95%CI：{klo:.2f}~{khi:.2f}）"))

# ---------------------------------------------------------------------------
# 附表S5 模型评价重估
# ---------------------------------------------------------------------------
se_app = hanley_mcneil_se(lg["apparent_auc"], 72, 31)
se_cor = hanley_mcneil_se(lg["corrected_auc"], 72, 31)
hdr = ["指标", "原稿报告", "重新评估", "说明"]
rows = [
    ["事件数与候选变量数之比（EPV）", "未报告", f"{epv:.1f}", "未改善组31例/8个候选变量；一般要求≥10"],
    ["表观AUC（95%CI）", f"{lg['apparent_auc']:.3f}", f"{lg['apparent_auc']:.3f}（{lg['apparent_auc'] - Z * se_app:.3f}~{lg['apparent_auc'] + Z * se_app:.3f}）", "Hanley-McNeil法"],
    ["校正AUC（95%CI）", f"{lg['corrected_auc']:.3f}（{lg['reported_auc_ci'][0]:.3f}~{lg['reported_auc_ci'][1]:.3f}）",
     f"{lg['corrected_auc']:.3f}（{lg['corrected_auc'] - Z * se_cor:.3f}~{lg['corrected_auc'] + Z * se_cor:.3f}）",
     "原稿区间为Bootstrap乐观度的波动范围，不是AUC的置信区间"],
    ["校准斜率", f"{lg['calibration_slope']:.2f}", "不适用", "在建模数据上计算的Logistic回归表观校准斜率恒为1"],
    ["预测变量时间顺序", "—", "不满足", "减重幅度在治疗结束时才能获得，不能用于治疗前预测"],
]
tables.append(write_table("附表S5_原列线图模型评价指标重估", hdr, rows, "注：基于原稿报告值计算；修改稿据此删除列线图"))

# ---------------------------------------------------------------------------
# 附表S6 样本量
# ---------------------------------------------------------------------------
def n_prec(p, w):
    return math.ceil(Z ** 2 * p * (1 - p) / w ** 2)


hdr = ["研究目标", "假设", "所需例数"]
n_sp = n_prec(0.30, 0.10)
n_se = n_prec(0.90, 0.05)
rows = [
    ["估计特异度", "预期30%，95%CI半宽±10%", f"非GERD患者{n_sp}例；按患病率{100 * pre:.1f}%，共约{math.ceil(n_sp / (1 - pre))}例"],
    ["估计灵敏度", "预期90%，95%CI半宽±5%", f"GERD患者{n_se}例；共约{math.ceil(n_se / pre)}例"],
    ["疗效模型外部验证", "结局事件与非事件各≥100例", f"按未改善率{100 * 31 / N:.1f}%，约{math.ceil(100 / (31 / N))}例"],
]
tables.append(write_table("附表S6_后续研究样本量估算", hdr, rows, "注：按正态近似公式 n=1.96²×p×(1−p)/d² 计算"))

# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
def plain(s):
    return (str(s).replace("*", "").replace("$\\bar{x}$", "x̄").replace("\\", "").replace("^a^", "a").replace("^b^", "b"))


wb = Workbook()
wb.remove(wb.active)
thin = Side(style="thin")
for tb in tables:
    ws = wb.create_sheet(tb["name"][:31])
    ws.append([plain(h) for h in tb["header"]])
    for row in tb["rows"]:
        ws.append([plain(c) for c in row])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.border = Border(top=thin, bottom=thin)
    for c in ws[ws.max_row]:
        c.border = Border(bottom=thin)
    if tb["note"]:
        ws.append([])
        ws.append([plain(tb["note"])])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=False)
    for col in ws.columns:
        width = max(len(plain(c.value or "")) for c in list(col)[: len(tb["rows"]) + 1])
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.6), 60)
ws = wb.create_sheet("附表S7_一致性核查")
with open(CHK / "一致性核查.csv", encoding="utf-8-sig") as fh:
    for row in csv.reader(fh):
        ws.append(row)
for c in ws[1]:
    c.font = Font(bold=True)
for col, w in zip("ABCDEF", (22, 34, 12, 12, 8, 50)):
    ws.column_dimensions[col].width = w
wb.save(TAB / "统计表汇总.xlsx")
print(f"[02] 生成 {len(tables)} 张表 + 一致性核查，已写入 {TAB}")
